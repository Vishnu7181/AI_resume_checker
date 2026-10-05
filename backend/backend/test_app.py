"""
Auto Test Suite — 10 Agent Testers
Run: pytest test_app.py -v   (backend must be running on :8000)
"""
import io
import requests
import pytest
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document

BASE = "http://127.0.0.1:8000"
JD = "Looking for Python developer with FastAPI, SQL, REST API and Docker experience"


# ---------- file generators ----------
def make_pdf():
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    text = c.beginText(50, 750)
    for line in [
        "Vishnu Kumar", "Email: vishnu@example.com",
        "Skills: Python, FastAPI, SQL, REST API, Docker, Git",
        "Experience:", "Built REST APIs with FastAPI serving 10k users",
        "Designed SQL databases and optimized queries",
        "Deployed applications using Docker",
    ]:
        text.textLine(line)
    c.drawText(text)
    c.save()
    return buf.getvalue()

def make_docx():
    doc = Document()
    doc.add_heading("Vishnu Kumar", 0)
    doc.add_paragraph("Skills: Python, FastAPI, SQL, REST API, Docker")
    doc.add_paragraph("Built REST APIs with FastAPI and SQL databases")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def make_fake_pdf():  # DOCX bytes renamed to .pdf
    return make_docx()

def make_txt():
    return b"Vishnu Kumar\nSkills: Python, FastAPI, SQL, REST API, Docker\n"

def make_empty():
    return b""

def upload(files=None, data=None):
    return requests.post(
        f"{BASE}/analyze-file",
        files=files or {},
        data=data or {"job_description": JD},
        timeout=30,
    )


# ---------- Agent Testers ----------
def test_1_backend_running():
    """Agent 0: backend is alive"""
    r = requests.get(f"{BASE}/health", timeout=5)
    assert r.status_code in (200, 404), "Backend not reachable"

def test_2_valid_pdf():
    """Agent 1: Upload Tester — valid text PDF returns 200 + scores"""
    r = upload(files={"file": ("resume.pdf", make_pdf(), "application/pdf")})
    assert r.status_code == 200, f"Got {r.status_code}: {r.text}"
    body = r.json()
    assert "score" in body or "match_score" in body, f"No score in response: {body}"

def test_3_valid_docx():
    """Agent 2: Format Tester — DOCX parses correctly"""
    r = upload(files={"file": ("resume.docx", make_docx(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, f"Got {r.status_code}: {r.text}"

def test_4_valid_txt():
    """Agent 3: TXT Tester — plain text parses"""
    r = upload(files={"file": ("resume.txt", make_txt(), "text/plain")})
    assert r.status_code == 200, f"Got {r.status_code}: {r.text}"

def test_5_fake_pdf_rejected():
    """Agent 4: Fake-PDF Tester — renamed DOCX rejected with clear message"""
    r = upload(files={"file": ("fake.pdf", make_fake_pdf(), "application/pdf")})
    assert r.status_code == 400, f"Expected 400, got {r.status_code}"
    assert "detail" in r.json(), f"No clear message: {r.text}"

def test_6_empty_file_rejected():
    """Agent 7: Empty Tester — 0-byte file rejected cleanly"""
    r = upload(files={"file": ("empty.pdf", make_empty(), "application/pdf")})
    assert r.status_code in (400, 422), f"Expected 400/422, got {r.status_code}"

def test_7_big_file_rejected():
    """Agent 6: Size Tester — >5MB rejected (or handled)"""
    big = b"%PDF-1.4\n" + b"x" * (6 * 1024 * 1024)
    r = upload(files={"file": ("big.pdf", big, "application/pdf")})
    assert r.status_code in (400, 413, 422), f"Got {r.status_code}"

def test_8_scores_change_with_input():
    """Agent 8: Score Tester — different inputs give different scores"""
    r1 = upload(files={"file": ("a.pdf", make_pdf(), "application/pdf")})
    assert r1.status_code == 200, r1.text
    b1 = r1.json()
    s1 = b1.get("score") or b1.get("match_score")
    assert isinstance(s1, (int, float)), f"Score not numeric: {b1}"
    # same file + totally different JD should give a different result
    r2 = requests.post(f"{BASE}/analyze-file",
        files={"file": ("a.pdf", make_pdf(), "application/pdf")},
        data={"job_description": "Sales representative, cold calling, CRM"},
        timeout=30)
    if r2.status_code == 200:
        b2 = r2.json()
        s2 = b2.get("score") or b2.get("match_score")
        assert s1 != s2, "Score identical for completely different JDs"

def test_9_text_mode_analysis():
    """Agent 8b: text paste mode works"""
    r = requests.post(f"{BASE}/analyze",
        data={"job_description": JD,
              "resume_text": "Skills: Python, FastAPI, SQL, Docker"},
        timeout=30)
    assert r.status_code == 200, f"Got {r.status_code}: {r.text}"

def test_10_cors_headers():
    """Agent 9: CORS configured for frontend :5173"""
    r = requests.options(f"{BASE}/analyze-file",
        headers={"Origin": "http://localhost:5173",
                 "Access-Control-Request-Method": "POST"}, timeout=5)
    assert r.headers.get("access-control-allow-origin") in (
        "http://localhost:5173", "http://127.0.0.1:5173", "*"), \
        f"CORS missing: {dict(r.headers)}"
